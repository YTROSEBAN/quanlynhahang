from fastapi import APIRouter, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import text
from database.db import engine
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO

router = APIRouter()
templates = Jinja2Templates(directory="templates")


@router.get("/hoadon", response_class=HTMLResponse)
def list_hoadon(request: Request):
    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT h.*, b.TenBan, n.HoTen, k.HoTen AS TenKH
                FROM hoadon h 
                LEFT JOIN ban b ON h.MaBan = b.MaBan 
                LEFT JOIN nhanvien n ON h.MaNV = n.MaNV
                LEFT JOIN khachhang k ON h.MaKH = k.MaKH
                ORDER BY h.MaHD DESC
            """))
            items = result.fetchall()
            bans = conn.execute(text("SELECT * FROM ban")).fetchall()
            nhanviens = conn.execute(text("SELECT * FROM nhanvien")).fetchall()
            monans = conn.execute(text("SELECT * FROM monan")).fetchall()
        return templates.TemplateResponse(request, "hoadon.html", {"items": items, "bans": bans, "nhanviens": nhanviens, "monans": monans, "active_page": "hoadon"})
    except Exception as e:
        return templates.TemplateResponse(request, "hoadon.html", {"error": str(e), "items": [], "bans": [], "nhanviens": [], "monans": [], "active_page": "hoadon"})


@router.post("/hoadon/add")
def add_hoadon(
    MaBan: int = Form(...), MaNV: int = Form(...),
    MaMon: list = Form([]), SoLuong: list = Form([])
):
    try:
        tong = 0
        with engine.connect() as conn:
            conn.execute(
                text("INSERT INTO hoadon (MaNV, MaBan, TongTien, TrangThai) VALUES (:manv, :maban, 0, 'Dang phuc vu')"),
                {"manv": MaNV, "maban": MaBan}
            )
            ma_hd = conn.execute(text("SELECT LAST_INSERT_ID()")).scalar()
            for i in range(len(MaMon)):
                mon = conn.execute(text("SELECT * FROM monan WHERE MaMon=:id"), {"id": MaMon[i]}).fetchone()
                sl = int(SoLuong[i]) if i < len(SoLuong) else 1
                if mon:
                    thanh_tien = mon.DonGia * sl
                    tong += thanh_tien
                    conn.execute(
                        text("INSERT INTO chitiethoadon (MaHD, MaMon, SoLuong, DonGia, ThanhTien) VALUES (:hd, :mon, :sl, :dg, :tt)"),
                        {"hd": ma_hd, "mon": MaMon[i], "sl": sl, "dg": mon.DonGia, "tt": thanh_tien}
                    )
            conn.execute(
                text("UPDATE hoadon SET TongTien=:tong WHERE MaHD=:id"),
                {"tong": tong, "id": ma_hd}
            )
            conn.execute(text("UPDATE ban SET TrangThai='Da dat' WHERE MaBan=:id"), {"id": MaBan})
            conn.commit()
    except Exception:
        pass
    return RedirectResponse(url="/hoadon", status_code=303)


@router.get("/hoadon/delete/{id}")
def delete_hoadon(id: int):
    try:
        with engine.connect() as conn:
            hd = conn.execute(text("SELECT MaBan FROM hoadon WHERE MaHD=:id"), {"id": id}).fetchone()
            if hd and hd.MaBan:
                conn.execute(text("UPDATE ban SET TrangThai='Trong' WHERE MaBan=:id"), {"id": hd.MaBan})
            conn.execute(text("DELETE FROM chitiethoadon WHERE MaHD = :id"), {"id": id})
            conn.execute(text("DELETE FROM hoadon WHERE MaHD = :id"), {"id": id})
            conn.commit()
    except Exception:
        pass
    return RedirectResponse(url="/hoadon", status_code=303)


@router.get("/hoadon/thanhtoan/{id}")
def thanh_toan(id: int):
    try:
        with engine.connect() as conn:
            hd = conn.execute(text("SELECT MaBan FROM hoadon WHERE MaHD=:id"), {"id": id}).fetchone()
            if hd and hd.MaBan:
                conn.execute(text("UPDATE ban SET TrangThai='Trong' WHERE MaBan=:id"), {"id": hd.MaBan})
            conn.execute(text("UPDATE hoadon SET TrangThai='Da thanh toan' WHERE MaHD=:id"), {"id": id})
            conn.commit()
    except Exception:
        pass
    return RedirectResponse(url="/hoadon", status_code=303)


@router.get("/hoadon/xuat/{id}", response_class=HTMLResponse)
def xuat_hoadon(request: Request, id: int):
    try:
        with engine.connect() as conn:
            hd = conn.execute(text("""
                SELECT h.*, b.TenBan, n.HoTen AS TenNV, k.HoTen AS TenKH
                FROM hoadon h
                LEFT JOIN ban b ON h.MaBan = b.MaBan
                LEFT JOIN nhanvien n ON h.MaNV = n.MaNV
                LEFT JOIN khachhang k ON h.MaKH = k.MaKH
                WHERE h.MaHD = :id
            """), {"id": id}).fetchone()
            if hd is None:
                return RedirectResponse(url="/hoadon", status_code=303)
            chitiet = conn.execute(text("""
                SELECT ct.*, m.TenMon
                FROM chitiethoadon ct
                LEFT JOIN monan m ON ct.MaMon = m.MaMon
                WHERE ct.MaHD = :id
            """), {"id": id}).fetchall()
        return templates.TemplateResponse(request, "hoadon_in.html", {"hd": hd, "chitiet": chitiet})
    except Exception:
        return RedirectResponse(url="/hoadon", status_code=303)


@router.get("/hoadon/excel/{id}")
def xuat_excel(id: int):
    try:
        with engine.connect() as conn:
            hd = conn.execute(text("""
                SELECT h.*, b.TenBan, n.HoTen AS TenNV, k.HoTen AS TenKH
                FROM hoadon h
                LEFT JOIN ban b ON h.MaBan = b.MaBan
                LEFT JOIN nhanvien n ON h.MaNV = n.MaNV
                LEFT JOIN khachhang k ON h.MaKH = k.MaKH
                WHERE h.MaHD = :id
            """), {"id": id}).fetchone()
            chitiet = conn.execute(text("""
                SELECT ct.*, m.TenMon
                FROM chitiethoadon ct
                LEFT JOIN monan m ON ct.MaMon = m.MaMon
                WHERE ct.MaHD = :id
            """), {"id": id}).fetchall()

        wb = Workbook()
        ws = wb.active
        ws.title = f"HoaDon_{id}"

        thin = Side(style='thin', color='CCCCCC')
        medium = Side(style='medium', color='2F5496')
        border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
        border_header = Border(left=thin, right=thin, top=medium, bottom=medium)
        border_bottom = Border(bottom=medium)

        center = Alignment(horizontal='center', vertical='center')
        right = Alignment(horizontal='right', vertical='center')
        left = Alignment(horizontal='left', vertical='center')
        wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)

        fill_title = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
        fill_header = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        fill_alt = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
        fill_total = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        fill_success = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')

        font_title = Font(bold=True, size=18, color='FFFFFF')
        font_subtitle = Font(bold=True, size=11, color='8DB4E2')
        font_bold = Font(bold=True, size=11)
        font_normal = Font(size=11)
        font_small = Font(size=10, color='666666')
        font_header_white = Font(bold=True, size=11, color='FFFFFF')
        font_total = Font(bold=True, size=13, color='C00000')
        font_money = Font(bold=True, size=12, color='2F5496')

        ws.column_dimensions['A'].width = 4
        ws.column_dimensions['B'].width = 32
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 22

        ws.merge_cells('A1:E1')
        title_cell = ws['A1']
        title_cell.value = 'NHÀ HÀNG'
        title_cell.font = font_title
        title_cell.fill = fill_title
        title_cell.alignment = center
        ws.row_dimensions[1].height = 40

        ws.merge_cells('A2:E2')
        sub = ws['A2']
        sub.value = 'HÓA ĐƠN BÁN HÀNG'
        sub.font = font_subtitle
        sub.fill = fill_title
        sub.alignment = center
        ws.row_dimensions[2].height = 25

        for col in range(1, 6):
            ws.cell(row=1, column=col).fill = fill_title
            ws.cell(row=2, column=col).fill = fill_title

        row = 4
        info_data = [
            ('📋 Số HĐ:', f'#{hd.MaHD}'),
            ('📅 Ngày lập:', hd.NgayLap.strftime('%d/%m/%Y %H:%M') if hd.NgayLap else 'N/A'),
            ('🪑 Bàn:', hd.TenBan or 'N/A'),
            ('👤 Nhân viên:', hd.TenNV or 'N/A'),
            ('🏠 Khách hàng:', hd.TenKH or 'Khách lẻ'),
        ]
        for label, val in info_data:
            c1 = ws.cell(row=row, column=1, value=label)
            c1.font = font_bold
            c2 = ws.cell(row=row, column=2, value=val)
            c2.font = font_normal
            row += 1

        row += 1
        ws.merge_cells(f'A{row}:E{row}')
        sep = ws.cell(row=row, column=1)
        sep.border = border_bottom
        for c in range(1, 6):
            ws.cell(row=row, column=c).border = Border(bottom=Side(style='thin', color='4472C4'))
        row += 1

        headers = ['STT', 'MÓN ĂN', 'SỐ LƯỢNG', 'ĐƠN GIÁ (₫)', 'THÀNH TIỀN (₫)']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = font_header_white
            cell.fill = fill_header
            cell.border = border_header
            cell.alignment = center
        ws.row_dimensions[row].height = 28
        row += 1

        for i, c in enumerate(chitiet, 1):
            is_alt = i % 2 == 0
            row_fill = fill_alt if is_alt else None

            ws.cell(row=row, column=1, value=i).border = border_all
            ws.cell(row=row, column=1).alignment = center
            ws.cell(row=row, column=1).font = font_bold
            if row_fill: ws.cell(row=row, column=1).fill = row_fill

            ws.cell(row=row, column=2, value=c.TenMon).border = border_all
            ws.cell(row=row, column=2).font = font_normal
            if row_fill: ws.cell(row=row, column=2).fill = row_fill

            ws.cell(row=row, column=3, value=c.SoLuong).border = border_all
            ws.cell(row=row, column=3).alignment = center
            ws.cell(row=row, column=3).font = font_normal
            if row_fill: ws.cell(row=row, column=3).fill = row_fill

            dg = ws.cell(row=row, column=4, value=float(c.DonGia))
            dg.border = border_all
            dg.number_format = '#,##0'
            dg.alignment = right
            dg.font = font_normal
            if row_fill: dg.fill = row_fill

            tt = ws.cell(row=row, column=5, value=float(c.ThanhTien))
            tt.border = border_all
            tt.number_format = '#,##0'
            tt.alignment = right
            tt.font = Font(bold=True, size=11, color='C00000' if float(c.ThanhTien) > 500000 else '333333')
            if row_fill: tt.fill = row_fill

            ws.row_dimensions[row].height = 24
            row += 1

        row += 1
        ws.merge_cells(f'A{row}:C{row}')

        total_label = ws.cell(row=row, column=4, value='TỔNG CỘNG:')
        total_label.font = Font(bold=True, size=12)
        total_label.alignment = right
        total_label.fill = fill_total

        total_val = ws.cell(row=row, column=5, value=float(hd.TongTien))
        total_val.font = font_total
        total_val.number_format = '#,##0 ₫'
        total_val.alignment = right
        total_val.fill = fill_total

        for c in range(4, 6):
            ws.cell(row=row, column=c).border = Border(
                top=Side(style='medium', color='C00000'),
                bottom=Side(style='double', color='C00000')
            )
        ws.row_dimensions[row].height = 30

        if hd.GiamGia and float(hd.GiamGia) > 0:
            row += 1
            ws.cell(row=row, column=4, value='Giảm giá:').font = font_bold
            ws.cell(row=row, column=4).alignment = right
            ws.cell(row=row, column=5, value=-float(hd.GiamGia)).font = Font(bold=True, size=11, color='FF0000')
            ws.cell(row=row, column=5).number_format = '#,##0 ₫'
            ws.cell(row=row, column=5).alignment = right

            row += 1
            ws.cell(row=row, column=4, value='THANH TOÁN:').font = Font(bold=True, size=13, color='006100')
            ws.cell(row=row, column=4).alignment = right
            ws.cell(row=row, column=4).fill = fill_success
            ws.cell(row=row, column=5, value=float(hd.ThanhTien or hd.TongTien))
            ws.cell(row=row, column=5).font = Font(bold=True, size=14, color='006100')
            ws.cell(row=row, column=5).number_format = '#,##0 ₫'
            ws.cell(row=row, column=5).alignment = right
            ws.cell(row=row, column=5).fill = fill_success
            for c in range(4, 6):
                ws.cell(row=row, column=c).border = Border(
                    top=Side(style='medium', color='006100'),
                    bottom=Side(style='double', color='006100')
                )

        row += 2
        tt_text = 'ĐÃ THANH TOÁN' if hd.TrangThai in ('Da thanh toan', 'Đã thanh toán') else 'CHƯA THANH TOÁN'
        tt_color = '006100' if 'ĐÃ' in tt_text.upper() or 'DA' in tt_text.upper().replace('Đ','D') else 'C00000'

        ws.merge_cells(f'A{row}:E{row}')
        status = ws.cell(row=row, column=1, value=f'Trạng thái: {tt_text}')
        status.font = Font(bold=True, size=12, color=tt_color)
        status.alignment = center

        row += 2
        ws.merge_cells(f'A{row}:E{row}')
        ws.cell(row=row, column=1, value='Cảm ơn quý khách! Hẹn gặp lại!').font = Font(italic=True, size=11, color='666666')
        ws.cell(row=row, column=1).alignment = center

        ws.print_area = f'A1:E{row}'
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = ws.PAPERSIZE_A5
        ws.page_margins.left = 0.5
        ws.page_margins.right = 0.5

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"HoaDon_{id}.xlsx"
        return StreamingResponse(
            buf,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        return RedirectResponse(url="/hoadon", status_code=303)


@router.get("/api/hoadon")
def api_hoadon():
    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT h.*, b.TenBan, n.HoTen AS TenNV
                FROM hoadon h 
                LEFT JOIN ban b ON h.MaBan = b.MaBan 
                LEFT JOIN nhanvien n ON h.MaNV = n.MaNV 
                ORDER BY h.MaHD DESC
            """))
            items = [dict(row._mapping) for row in result]
        return items
    except Exception as e:
        return {"error": str(e), "items": []}


@router.get("/api/chitiet/{ma_hd}")
def api_chitiet(ma_hd: int):
    try:
        with engine.connect() as conn:
            result = conn.execute(text("""
                SELECT c.*, m.TenMon 
                FROM chitiethoadon c 
                LEFT JOIN monan m ON c.MaMon = m.MaMon 
                WHERE c.MaHD = :id
            """), {"id": ma_hd})
            items = [dict(row._mapping) for row in result]
        return items
    except Exception as e:
        return {"error": str(e), "items": []}
